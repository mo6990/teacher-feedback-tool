"""教师反馈完成率看板。运行：streamlit run feedback_tool.py"""

from io import BytesIO
from pathlib import Path
from html import escape
import math
import re
import unicodedata

import pandas as pd
import streamlit as st
from PIL import Image, ImageDraw, ImageFont


DEFAULT_NAMES = "彭钰茂、徐小涵、高弋力、杨安小、顾毅清云、姚蓓蓓、张洊维、赵万铖、陈祉璇、李丽旋、易诗樵、蔡天翔、王皓玥、孙天玥、高子力、胡海枫、张宇瀚、宋蕊君、魏志如、潘蓉、吴蕊、夏梦、刘芳菲、王诗仪、王馨悦、王晨亚盈、刘悦、吴恙、王千羽、吴丹妮、任雨铃、李睿"
RATE = "班级反馈完成率"
BASE_DIR = Path(__file__).resolve().parent
LEVEL_COLORS = {
    "优秀": ("#047857", "#DCFCE7"),
    "合格": ("#1D4ED8", "#DBEAFE"),
    "待改进": ("#B45309", "#FEF3C7"),
}


def parse_names(value):
    return list(dict.fromkeys(n.strip() for n in re.split(r"[、,，;；\n]+", value) if n.strip()))


def get_level(rate):
    return "优秀" if rate >= 0.9 else "合格" if rate >= 0.7 else "待改进"


def calculate_teacher_stats(df, target_names):
    """保留原考核口径：一行视为一课次，仅 Y 算完成。"""
    data = df.copy()
    data.columns = data.columns.astype(str).str.strip()
    if data.columns.duplicated().any():
        raise ValueError("存在重复列名，请先检查 Excel 表头。")
    project_col = next((col for col in ("项目", "产品品类") if col in data), None)
    required = ["教师姓名", "是否进行班级学情反馈"]
    missing = [col for col in required if col not in data]
    if project_col is None:
        missing.append("项目（或产品品类）")
    if missing:
        raise ValueError("缺少字段：" + "、".join(missing))
    data = data.dropna(subset=["教师姓名", project_col])
    data["教师姓名"] = data["教师姓名"].astype(str).str.strip()
    data[project_col] = data[project_col].astype(str).str.strip()
    data = data[(data[project_col] != project_col) & (data[project_col] != "")]
    data = data[data["教师姓名"].isin(target_names)].reset_index(drop=True)
    if data.empty:
        return pd.DataFrame(), pd.DataFrame(), data
    data["是否进行班级学情反馈"] = (
        data["是否进行班级学情反馈"].fillna("").astype(str).str.strip().str.upper()
    )
    data["班级反馈完成"] = data["是否进行班级学情反馈"].eq("Y")
    stats = data.groupby("教师姓名", sort=False).agg(
        总授课课次=("教师姓名", "size"),
        班级反馈完成课次=("班级反馈完成", "sum"),
        负责项目=(project_col, lambda x: "、".join(sorted(set(x)))),
    ).reset_index()
    if "科目名称" in data:
        subjects = data.groupby("教师姓名")["科目名称"].apply(
            lambda x: "、".join(sorted({str(v).strip() for v in x.dropna() if str(v).strip()}))
        )
        stats["负责科目"] = stats["教师姓名"].map(subjects)
    stats[RATE] = stats["班级反馈完成课次"] / stats["总授课课次"]
    stats["综合完成率"] = stats[RATE]
    stats["考核等级"] = stats[RATE].map(get_level)
    stats = stats.sort_values(RATE, ascending=False, kind="stable").reset_index(drop=True)
    show = stats.copy()
    for col in (RATE, "综合完成率"):
        show[col] = show[col].map(lambda x: f"{x:.1%}")
    return stats, show, data


def get_font_path():
    # 云端优先使用随项目提交的字体，不依赖运行时联网下载。
    candidates = [
        BASE_DIR / "fonts" / "NotoSansSC-Regular.otf",
        BASE_DIR / "NotoSansSC-Regular.otf",
        BASE_DIR / "SimHei.ttf",
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    ]
    for candidate in candidates:
        if candidate.is_file():
            return str(candidate)
    raise FileNotFoundError("缺少中文字体。请一并上传 fonts 文件夹内的 NotoSansSC-Regular.otf。")


@st.cache_data(show_spinner=False, max_entries=16)
def generate_chart_image(normal_stats, vip_stats, target_names, title_text, subtitle_text, compact=True):
    """按像素排版并以 2 倍分辨率绘制，避免比例坐标导致文字重叠。"""
    scale, width = 2, 1440
    font_path = get_font_path()
    fonts = {size: ImageFont.truetype(font_path, size * scale) for size in (14, 16, 18, 20, 22, 30, 32)}
    scratch = ImageDraw.Draw(Image.new("RGB", (10, 10)))

    def text_width(text, size):
        return scratch.textlength(str(text), font=fonts[size]) / scale

    def wrap(text, size, max_width):
        lines = []
        for paragraph in str(text).split("\n"):
            line = ""
            for char in paragraph:
                if line and text_width(line + char, size) > max_width:
                    lines.append(line)
                    line = char
                else:
                    line += char
            lines.append(line)
        return lines or [""]

    def lookup(frame):
        if frame is None or frame.empty:
            return {}
        return {row["教师姓名"]: row for row in frame.to_dict("records")}

    normal, vip = lookup(normal_stats), lookup(vip_stats)
    rows = [(name, normal.get(name), vip.get(name)) for name in dict.fromkeys(target_names)]
    if not rows:
        raise ValueError("请至少填写一位教师。")

    def rank(row):
        rates = [record[RATE] for record in row[1:] if record is not None]
        return sum(rates) / len(rates) if rates else -1

    rows.sort(key=rank, reverse=True)
    title_lines = wrap(title_text.strip() or "教师学情反馈完成率", 32, 1288)
    subtitle_lines = wrap(subtitle_text.strip(), 18, 1288) if subtitle_text.strip() else []
    title_y = 28 if compact else 86
    title_step = 42 if compact else 48
    subtitle_y = title_y + len(title_lines) * title_step + (8 if compact else 12)
    cards_y = subtitle_y + len(subtitle_lines) * 29 + (18 if compact else 30)
    table_y = cards_y if compact else cards_y + 140
    header_h, row_min, name_step = (60, 64, 28) if compact else (82, 104, 31)
    row_heights = [max(row_min, len(wrap(row[0], 22, 208)) * name_step + (20 if compact else 36)) for row in rows]
    footer_gap = 16 if compact else 24
    footer_y = table_y + header_h + sum(row_heights) + footer_gap
    height = footer_y + (72 if compact else 112)
    image = Image.new("RGB", (width * scale, height * scale), "#F3F6FA")
    draw = ImageDraw.Draw(image)

    def rect(box, fill, radius=0, outline=None):
        coords = tuple(round(v * scale) for v in box)
        draw.rounded_rectangle(coords, radius=round(radius * scale), fill=fill, outline=outline, width=scale)

    def line(points, color, thickness=1):
        draw.line([(round(x * scale), round(y * scale)) for x, y in points], fill=color, width=thickness * scale)

    def text(x, y, value, size=18, color="#172B4D", align="left"):
        value = str(value)
        if align == "right":
            x -= text_width(value, size)
        draw.text((round(x * scale), round(y * scale)), value, font=fonts[size], fill=color, anchor="lt")

    if not compact:
        rect((48, 44, 54, 64), "#2563EB", 3)
        text(68, 45, "教学运营  /  学情反馈", 16, "#64748B")
    for i, title in enumerate(title_lines):
        text(48, title_y + title_step * i, title, 32)
    for i, subtitle in enumerate(subtitle_lines):
        text(48, subtitle_y + 29 * i, subtitle, 18, "#64748B")

    available = set(normal) | set(vip)
    needs_attention = {name for name, a, b in rows if any(r is not None and r[RATE] < 0.7 for r in (a, b))}

    def average(records):
        vals = [record[RATE] for name, record in records.items() if name in target_names]
        return f"{sum(vals) / len(vals):.1%}" if vals else "—"

    cards = [
        ("有课教师", f"{len(available & set(target_names))} 位", "#172B4D"),
        ("班课 · 教师平均完成率", average(normal), "#2563EB"),
        ("VIP · 教师平均完成率", average(vip), "#059669"),
        ("待跟进教师 · 已去重", f"{len(needs_attention)} 位", "#C27816"),
    ]
    if not compact:
        for i, (label, value, color) in enumerate(cards):
            x = 48 + i * 340
            rect((x, cards_y, x + 324, cards_y + 112), "white", 16)
            text(x + 22, cards_y + 20, label, 16, "#64748B")
            text(x + 22, cards_y + 52, value, 32, color)

    table_end = footer_y - footer_gap
    rect((48, table_y, 1392, table_end), "white", 18)
    heading_y = table_y + (16 if compact else 25)
    text(76, heading_y, "教师姓名", 20)
    rect((330, heading_y + 4, 338, heading_y + 24), "#2563EB", 4)
    text(352, heading_y, "班课学情反馈率", 20, "#1D4ED8")
    rect((866, heading_y + 4, 874, heading_y + 24), "#059669", 4)
    text(888, heading_y, "VIP 学情反馈率", 20, "#047857")
    line([(76, table_y + header_h - 10), (1364, table_y + header_h - 10)], "#E2E8F0")

    def draw_rate(x, y, record, color, track_color, empty_label):
        right = x + 462
        track_y = y + (42 if compact else 49)
        track_h = 8 if compact else 12
        if record is None:
            text(x, y + 8, empty_label, 18, "#94A3B8")
            rect((x, track_y, right, track_y + track_h), "#F1F5F9", track_h / 2)
            return
        rate = float(record[RATE])
        level = get_level(rate)
        text(x, y + 8, f"{rate:.2%}", 22, color)
        rect((x, track_y, right, track_y + track_h), track_color, track_h / 2)
        if rate > 0:
            fill_width = min(462, max(1, round(462 * rate)))
            rect((x, track_y, x + fill_width, track_y + track_h), color, min(track_h / 2, fill_width / 2))
        level_color = "#B7791F" if rate < 0.7 else "#64748B"
        if compact:
            text(right, y + 14, level, 16, level_color, "right")
        else:
            text(x, y + 70, level, 14, level_color)

    y = table_y + header_h
    for index, ((name, normal_record, vip_record), row_height) in enumerate(zip(rows, row_heights)):
        if index % 2 == 1:
            rect((64, y, 1376, y + row_height), "#F8FAFC", 10)
        name_lines = wrap(name, 22, 208)
        name_y = y + (row_height - len(name_lines) * name_step) / 2
        for j, name_line in enumerate(name_lines):
            text(76, name_y + j * name_step, name_line, 22)
        draw_rate(330, y + (row_height - row_min) / 2, normal_record, "#2563EB", "#E7EEFF", "暂无班课记录")
        draw_rate(866, y + (row_height - row_min) / 2, vip_record, "#059669", "#DFF5EB", "暂无 VIP 记录")
        y += row_height

    text(52, footer_y, "优秀 ≥ 90%     合格 ≥ 70% 且 < 90%     待改进 < 70%", 16, "#64748B")
    text(52, footer_y + 29, "完成率 = 班级反馈 Y 课次 / 总课次；无课记录不计入均值。", 14, "#94A3B8")
    if not compact:
        text(52, footer_y + 53, "按班课、VIP 非空完成率的算术平均降序排列；待跟进指任一班型完成率低于 70%。", 14, "#94A3B8")
    output = BytesIO()
    image.save(output, format="PNG", optimize=True, dpi=(240, 240))
    return output.getvalue()


@st.cache_data(show_spinner=False, max_entries=16)
def generate_excel(sheets):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.formatting.rule import DataBarRule
    from openpyxl.utils import get_column_letter

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, index=False, sheet_name=name)
            sheet = writer.sheets[name]
            accent = "059669" if "VIP" in name else "2563EB"
            pale = "F0FAF5" if "VIP" in name else "F3F7FE"
            sheet.sheet_properties.tabColor = accent
            sheet.sheet_view.showGridLines = False
            sheet.freeze_panes = "B2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.row_dimensions[1].height = 36
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.print_title_rows = "1:1"
            sheet.print_options.horizontalCentered = True
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor=accent)
                cell.font = Font(name="等线", size=11, color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            def display_width(value):
                return max((sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in part)
                            for part in str(value).splitlines()), default=0)

            for index, col in enumerate(frame.columns, 1):
                longest = max([display_width(col)] + [display_width(v) for v in frame[col].head(100)])
                sheet.column_dimensions[get_column_letter(index)].width = min(44, max(16, longest + 4))
                if col in (RATE, "综合完成率"):
                    for cells in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                        if isinstance(cells[0].value, (int, float)):
                            cells[0].number_format = "0.0%"
                    if len(frame):
                        letter = get_column_letter(index)
                        sheet.conditional_formatting.add(
                            f"{letter}2:{letter}{len(frame) + 1}",
                            DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                                        color=accent, showValue=True),
                        )
            for row in sheet.iter_rows(min_row=2):
                row_height = 30
                for cell in row:
                    # 原始表格中的文本原样导出，避免以 = 开头的文本变为公式。
                    if cell.data_type == "f":
                        cell.data_type = "s"
                    col = str(frame.columns[cell.column - 1])
                    cell.font = Font(name="等线", size=11, color="334155", bold=col == "教师姓名")
                    centered = isinstance(cell.value, (int, float, bool)) or col in ("考核等级", "是否进行班级学情反馈", "班型")
                    cell.alignment = Alignment(horizontal="center" if centered else "left", vertical="center", wrap_text=True)
                    cell.fill = PatternFill("solid", fgColor=pale if cell.row % 2 == 0 else "FFFFFF")
                    cell.border = Border(bottom=Side(style="hair", color="E2E8F0"))
                    if col == "考核等级" and cell.value in LEVEL_COLORS:
                        foreground, background = LEVEL_COLORS[cell.value]
                        cell.font = Font(name="等线", size=11, bold=True, color=foreground[1:])
                        cell.fill = PatternFill("solid", fgColor=background[1:])
                    elif col == "是否进行班级学情反馈":
                        is_complete = str(cell.value).upper() == "Y"
                        cell.font = Font(name="等线", size=11, bold=True, color="047857" if is_complete else "B45309")
                        cell.fill = PatternFill("solid", fgColor="DCFCE7" if is_complete else "FEF3C7")
                    width = sheet.column_dimensions[cell.column_letter].width - 3
                    wrapped_lines = sum(max(1, math.ceil(display_width(part) / width)) for part in str(cell.value or "").split("\n"))
                    row_height = max(row_height, 18 * wrapped_lines + 10)
                sheet.row_dimensions[row[0].row].height = min(120, row_height)
    return output.getvalue()


def result_table_html(frame, title, theme="blue"):
    """静态展示表：所有来自表格的文本均转义，数字进度单独绘制。"""
    accent, pale = ("#059669", "#ECFDF5") if theme == "green" else ("#2563EB", "#EFF6FF")
    css = """
    <style>
    .feedback-table-card {background:white;border:1px solid #E2E8F0;border-radius:12px;overflow:hidden;max-width:1100px;
        color:#334155;font-family:Inter,'Microsoft YaHei',sans-serif;margin:8px 0 12px;box-shadow:0 4px 20px #172b4d04;}
    .feedback-table-top {display:flex;align-items:center;justify-content:space-between;gap:12px;padding:12px 16px;line-height:22px;}
    .feedback-table-top strong {font-size:16px;color:#172B4D;min-width:0;flex:1;overflow-wrap:anywhere;}
    .feedback-table-top span {font-size:12px;color:#64748B;white-space:nowrap;flex-shrink:0;}
    .feedback-table-scroll {overflow:auto;max-height:560px;}
    .feedback-table {width:100%;border-collapse:collapse;font-size:14px;line-height:20px;white-space:nowrap;}
    .feedback-table thead th {background:var(--table-pale);color:var(--table-accent);font-weight:600;
        padding:9px 12px;text-align:left;border-bottom:1px solid #E2E8F0;position:sticky;top:0;z-index:1;}
    .feedback-table td {padding:7px 12px;border-bottom:1px solid #EDF1F6;vertical-align:middle;}
    .feedback-table tbody tr:nth-child(even) {background:#FAFBFD;}
    .feedback-table tbody tr:hover {background:var(--table-pale);}
    .feedback-table tr:last-child td {border-bottom:0;}
    .feedback-table .teacher-name {font-weight:600;color:#172B4D;}
    .feedback-table .number-cell {text-align:center;font-variant-numeric:tabular-nums;}
    .feedback-table .muted-cell {color:#64748B;max-width:200px;white-space:normal;min-width:64px;overflow-wrap:anywhere;}
    .feedback-table .order-cell {color:#94A3B8;text-align:center;font-size:13px;}
    .feedback-pill {display:inline-block;padding:2px 9px;border-radius:6px;font-size:12px;font-weight:600;}
    .feedback-rate {display:flex;align-items:center;gap:10px;min-width:150px;max-width:200px;}
    .feedback-rate-label {flex:0 0 56px;color:var(--table-accent);font-weight:600;font-variant-numeric:tabular-nums;}
    .feedback-track {flex:1;min-width:72px;height:5px;border-radius:4px;background:#E8EEF5;overflow:hidden;}
    .feedback-track-fill {height:100%;border-radius:4px;background:var(--table-accent);}
    .feedback-table-foot {padding:8px 16px;line-height:18px;background:#FAFBFD;border-top:1px solid #EDF1F6;font-size:12px;color:#64748B;}
    </style>
    """
    cols = list(frame.columns)
    labels = {"负责项目": "项目", "负责科目": "科目", "总授课课次": "总课次",
              "班级反馈完成课次": "已反馈课次", RATE: "反馈完成率"}
    headers = '<th class="order-cell">序号</th>'
    for col in cols:
        cls = ' class="number-cell"' if col in ("总授课课次", "班级反馈完成课次") else ""
        headers += f'<th{cls} title="{escape(str(col), quote=True)}">{escape(labels.get(col, str(col)))}</th>'
    rows = []
    for number, record in enumerate(frame.to_dict("records"), 1):
        cells = [f'<td class="order-cell">{number:02d}</td>']
        for col in cols:
            value = record[col]
            if pd.isna(value):
                cells.append('<td class="muted-cell">—</td>')
            elif col in (RATE, "综合完成率"):
                rate = float(value)
                percent = max(0, min(100, rate * 100))
                cells.append(f'<td><div class="feedback-rate"><div class="feedback-rate-label">{rate:.1%}</div>'
                             f'<div class="feedback-track"><div class="feedback-track-fill" style="width:{percent:.3f}%"></div></div></div></td>')
            elif col == "考核等级":
                fg, bg = LEVEL_COLORS.get(str(value), ("#64748B", "#F1F5F9"))
                cells.append(f'<td><span class="feedback-pill" style="color:{fg};background:{bg}">{escape(str(value))}</span></td>')
            elif col == "班型":
                fg, bg = ("#047857", "#DCFCE7") if "VIP" in str(value) else ("#1D4ED8", "#DBEAFE")
                cells.append(f'<td><span class="feedback-pill" style="color:{fg};background:{bg}">{escape(str(value))}</span></td>')
            else:
                cls = "teacher-name" if col == "教师姓名" else "number-cell" if isinstance(value, (int, float)) else "muted-cell"
                cells.append(f'<td class="{cls}">{escape(str(value))}</td>')
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return (css + f'<div class="feedback-table-card" style="--table-accent:{accent};--table-pale:{pale}">'
            f'<div class="feedback-table-top"><strong>{escape(title)}</strong><span>共 {len(frame)} 条记录</span></div>'
            f'<div class="feedback-table-scroll"><table class="feedback-table"><thead><tr>{headers}</tr></thead>'
            f'<tbody>{"".join(rows)}</tbody></table></div>'
            '<div class="feedback-table-foot">优秀 ≥ 90%　 ·　 合格 ≥ 70% 且 &lt; 90%　 ·　 待改进 &lt; 70%</div></div>')


def interactive_table(frame):
    """保留 Streamlit 原生表格的排序、搜索、复制和全屏查看功能。"""
    view = frame.reset_index(drop=True).copy()
    configs = {}
    for col in (RATE, "综合完成率"):
        if col in view:
            view[col] *= 100
            configs[col] = st.column_config.ProgressColumn(col, format="%.1f%%", min_value=0, max_value=100, width="medium")
    for col in ("总授课课次", "班级反馈完成课次"):
        if col in view:
            configs[col] = st.column_config.NumberColumn(col, format="%d 次", width="small")

    def row_style(row):
        background = "#F8FAFC" if row.name % 2 else "#FFFFFF"
        styles = []
        for col, value in row.items():
            style = f"background-color: {background}; color: #334155;"
            if col == "考核等级" and str(value) in LEVEL_COLORS:
                fg, bg = LEVEL_COLORS[str(value)]
                style = f"background-color: {bg}; color: {fg}; font-weight: bold;"
            elif col == "是否进行班级学情反馈":
                fg, bg = ("#047857", "#DCFCE7") if str(value) == "Y" else ("#B45309", "#FEF3C7")
                style = f"background-color: {bg}; color: {fg}; font-weight: bold;"
            styles.append(style)
        return styles

    # 大明细表使用原生渲染，避免逐单元格样式造成页面卡顿。
    if view.size <= 50000:
        display = view.style.apply(row_style, axis=1)
    else:
        display = view
    st.dataframe(display, hide_index=True, use_container_width=True, column_config=configs,
                 height=min(560, max(140, (len(view) + 1) * 36 + 4)))


def edit_table_title(default_title, key):
    # 独立保存标题，切换班型导致输入框暂时消失时，也不会丢失已编辑内容。
    saved_key = f"saved_table_title_{key}"
    widget_key = f"_table_title_{key}"
    if saved_key not in st.session_state:
        st.session_state[saved_key] = default_title
    st.session_state[widget_key] = st.session_state[saved_key]

    def save_title():
        st.session_state[saved_key] = st.session_state[widget_key]

    with st.expander("自定义表格标题", expanded=False):
        st.text_input("表格标题", key=widget_key, max_chars=160, on_change=save_title,
                      help="可填写组别、统计周期等信息；留空时使用默认标题。修改后按回车生效。")
    return st.session_state[saved_key].strip() or default_title


def show_result_table(frame, title, theme="blue", key="result"):
    if frame.empty:
        st.info("当前筛选条件下没有记录。")
        return
    custom_title = edit_table_title(title, key)
    st.html(result_table_html(frame, custom_title, theme))
    with st.expander("交互表格 · 排序、搜索与复制"):
        interactive_table(frame)


def show_ranking(frame, key, title="教师考核排名", theme="blue"):
    view = frame.copy()
    if "负责科目" in view:
        options = sorted({part for item in view["负责科目"].dropna() for part in str(item).split("、") if part})
        if options:
            selected = st.multiselect("按科目筛选", options, default=options, key=key)
            view = view[view["负责科目"].map(lambda value: any(part in selected for part in str(value).split("、")))]
    cols = [col for col in ("教师姓名", "负责项目", "负责科目", "总授课课次", "班级反馈完成课次", RATE, "考核等级") if col in view]
    show_result_table(view[cols], title, theme, key=key)


def main():
    st.set_page_config(page_title="学情反馈 · 教师考核", page_icon="📊", layout="wide")
    st.markdown("""
    <style>
    .stApp {background: #f5f7fb;}
    .block-container {max-width: 1440px; padding-top: 2.2rem; padding-bottom: 3rem;}
    [data-testid="stSidebar"] {background: #fff; border-right: 1px solid #e6ebf2;}
    h1, h2, h3 {color: #172b4d; letter-spacing: -.025em;}
    [data-testid="stMetric"] {background: white; border: 1px solid #e6ebf2;
        border-radius: 16px; padding: 18px 22px; min-height: 125px;}
    [data-testid="stMetricLabel"] {color: #64748b;}
    [data-testid="stMetricValue"] {color: #172b4d; font-weight: 650;}
    [data-testid="stFileUploader"] {background: white; border-radius: 14px; padding: 12px;}
    [data-testid="stTabs"] [role="tablist"] {gap: 20px;}
    [data-testid="stTabs"] [role="tab"] {font-size: 15px; padding-bottom: 12px;}
    [data-testid="stDownloadButton"] button {border-radius: 10px;}
    [data-testid="stDataFrame"] {border:1px solid #E2E8F0;border-radius:12px;overflow:hidden;}
    .page-kicker {color: #2563eb; font-size: 13px; font-weight: 700; letter-spacing: .12em; margin-bottom: 8px;}
    .page-note {color: #64748b; font-size: 15px; margin: -4px 0 24px; line-height: 1.8;}
    @media (max-width: 640px) {
        .block-container {padding: 1.2rem 1rem 2rem;}
        [data-testid="stMetric"] {padding: 14px; min-height: 110px;}
    }
    </style>
    """, unsafe_allow_html=True)
    with st.sidebar:
        st.subheader("考核设置")
        st.caption("编辑名单后，结果会自动更新。")
        name_input = st.text_area("教师名单", DEFAULT_NAMES, height=190,
                                  help="支持顿号、逗号、分号或换行分隔；重复姓名会自动去重。")
        target_names = parse_names(name_input)
        st.caption(f"当前名单 · {len(target_names)} 位教师")
        st.divider()
        st.markdown("**完成率考核规则**")
        st.markdown("🟢 **优秀**　≥ 90%\n\n🔵 **合格**　70%–90%（不含 90%）\n\n🟠 **待改进**　< 70%")
        st.caption("完成率 = 班级反馈 Y 课次 ÷ 总授课课次。普通班与 VIP 班独立考核。")
        with st.expander("查看上传要求"):
            st.markdown("上传 `.xlsx` 文件，第一行为表头，需包含：\n- 项目或产品品类\n- 教师姓名\n- 是否进行班级学情反馈")
            st.caption("支持只上传一个班型；仅统计名单内教师。每一行按一课次统计。空白反馈也计入总课次，但不算完成。")

    st.markdown('<div class="page-kicker">教学运营 / FEEDBACK INSIGHTS</div>', unsafe_allow_html=True)
    st.title("教师学情反馈看板")
    st.markdown('<div class="page-note">用一张图看清反馈进度，让每一次课后跟进都有据可依。</div>', unsafe_allow_html=True)
    left, right = st.columns(2)
    with left:
        normal_file = st.file_uploader("班课 · 上传普通班反馈表", type=["xlsx"], key="normal_upload")
    with right:
        vip_file = st.file_uploader("VIP · 上传 VIP 反馈表", type=["xlsx"], key="vip_upload")
    if not target_names:
        st.info("请在左侧填写至少一位教师，再开始统计。")
        return
    if normal_file is None and vip_file is None:
        st.info("上传至少一份反馈表，即可查看完成率、教师排名和待跟进清单。")
        return

    results, details = {}, {}
    for label, uploaded in (("普通班", normal_file), ("VIP班", vip_file)):
        if uploaded is None:
            continue
        try:
            frame = pd.read_excel(uploaded)
            stats, _, detail = calculate_teacher_stats(frame, target_names)
            if stats.empty:
                st.warning(f"{label}：未找到名单内教师的有效课次，请核对姓名和项目字段。")
                continue
            results[label], details[label] = stats, detail
            unexpected = ~detail["是否进行班级学情反馈"].isin(["Y", "N"])
            if unexpected.any():
                st.warning(f"{label}有 {int(unexpected.sum())} 条反馈值为空或非 Y/N，已按未完成计入，请核对原表。")
        except Exception as exc:
            st.error(f"{label}读取失败：{exc}")
    if not results:
        return

    all_stats = pd.concat([frame.assign(班型=label) for label, frame in results.items()], ignore_index=True)
    bad = all_stats[all_stats[RATE] < 0.7].copy()
    st.subheader("完成情况概览")
    cols = st.columns(4)
    cols[0].metric("有课教师", f"{all_stats['教师姓名'].nunique()} 位", help="至少一个班型有课次的教师人数，跨班型去重。")
    for col, label, text_label in ((cols[1], "普通班", "班课 · 教师平均完成率"), (cols[2], "VIP班", "VIP · 教师平均完成率")):
        col.metric(text_label, f"{results[label][RATE].mean():.1%}" if label in results else "—",
                   help="有课教师完成率的算术平均；不是全部课次的加权完成率。")
    cols[3].metric("待跟进教师", f"{bad['教师姓名'].nunique()} 位", help="任一班型完成率低于 70% 的教师人数，跨班型去重。")
    st.caption("按当前教师名单统计 · 缺少课次的班型显示为「—」，不当作 0%。")

    tab_chart, tab_ranking, tab_followup, tab_details = st.tabs(["反馈率对照图", "教师考核排名", "待跟进清单", "原始课次明细"])
    with tab_chart:
        with st.expander("自定义图表标题", expanded=False):
            chart_title = st.text_input("图表标题", "贵阳新东方国际教育雅思阅写组学情反馈", max_chars=160)
            chart_subtitle = st.text_input("副标题 / 统计周期", "班课与 VIP · 按教师逐项对照", max_chars=240)
        layout_col, width_col = st.columns(2)
        with layout_col:
            chart_layout = st.radio("图表布局", ["紧凑截图", "完整展示"], horizontal=True, key="chart_layout")
        with width_col:
            preview_width = st.slider("预览宽度", min_value=640, max_value=1200, value=960, step=80,
                                      help="调小后更方便整张截图；下载图片仍保持高清分辨率。")
        try:
            chart = generate_chart_image(results.get("普通班"), results.get("VIP班"), target_names, chart_title, chart_subtitle,
                                         compact=chart_layout == "紧凑截图")
            st.image(chart, width=preview_width)
            st.download_button("下载高清图表 PNG", chart, "学情反馈率图表.png", "image/png", type="primary")
            st.caption("按两个班型非空完成率的算术平均排序。下载原图可查看高清文字。")
        except Exception as exc:
            st.error(f"图表生成失败：{exc}")

    with tab_ranking:
        labels = list(results)
        if len(results) == 2:
            labels.append("综合汇总")
        selected = st.radio("查看班型", labels, horizontal=True, key="ranking_class")
        if selected == "综合汇总":
            combined = all_stats.groupby("教师姓名")[["总授课课次", "班级反馈完成课次"]].sum().reset_index()
            combined[RATE] = combined["班级反馈完成课次"] / combined["总授课课次"]
            combined["考核等级"] = combined[RATE].map(get_level)
            st.caption("综合完成率按两个班型的课次合并计算；各班型的独立考核结果保持不变。")
            show_ranking(combined.sort_values(RATE, ascending=False), "subjects_combined", "综合汇总 · 教师考核")
        else:
            show_ranking(results[selected], "subjects_" + selected, selected + " · 教师考核",
                         "green" if selected == "VIP班" else "blue")
    with tab_followup:
        if bad.empty:
            st.success("当前有课教师在各班型的完成率均达到 70%，暂无待跟进记录。")
        else:
            st.caption(f"{bad['教师姓名'].nunique()} 位教师 · {len(bad)} 条班型考核记录 · 按完成率从低到高排列")
            followup = bad.sort_values(RATE)[["班型", "教师姓名", "总授课课次", "班级反馈完成课次", RATE, "考核等级"]].copy()
            show_result_table(followup, "待跟进教师 · 优先关注", key="followup")
    with tab_details:
        selected_detail = st.radio("明细班型", list(details), horizontal=True, key="detail_class")
        detail_title = edit_table_title(selected_detail + " · 课次明细", "detail_" + selected_detail)
        st.subheader(detail_title)
        st.caption(f"仅展示名单内教师的有效课次，共 {len(details[selected_detail])} 条。")
        interactive_table(details[selected_detail])

    st.divider()
    sheets = {}
    for label, frame in results.items():
        sheets[label + "考核结果"] = frame
        sheets[label + "课次明细"] = details[label]
    st.download_button("导出完整考核结果 Excel", generate_excel(sheets), "教师反馈完成率考核结果.xlsx",
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    st.caption("包含已上传班型的考核结果和课次明细；页面科目筛选不影响完整导出。")


if __name__ == "__main__":
    main()
